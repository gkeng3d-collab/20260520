#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
คัดกรองหุ้น SET สำหรับเทรด จากไฟล์ listSET_YYYYMMDD.xlsx (ชีท ALL, 1 แถว = 1 หุ้น)

วิธีใช้:
    python3 screen_set_stocks.py <path ไฟล์ listSET_YYYYMMDD.xlsx> [ไฟล์ผลลัพธ์.xlsx]

กลยุทธ์ 3 แบบ:
  A) Momentum  — เทรนด์ขาขึ้นแข็งแรง โมเมนตัมดี ยังไม่ overbought
  B) Breakout  — เพิ่งทะลุฐาน/แนวต้าน หรือเพิ่งเกิดสัญญาณซื้อใหม่ พร้อมวอลุ่มยืนยัน
  C) Pullback  — ขาขึ้นใหญ่ (รายสัปดาห์) ยังดี แต่รายวันย่อลงมาใกล้แนวรับ EMA26

ผลลัพธ์: ไฟล์ Excel มีชีทเกณฑ์ + รายชื่อหุ้นแต่ละกลยุทธ์ เรียงตามคะแนนรวม
หมายเหตุ: เป็นการคัดกรองทางเทคนิคจากข้อมูลในไฟล์เท่านั้น ไม่ใช่คำแนะนำการลงทุน
"""
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------- เกณฑ์ Universe (สภาพคล่อง/คุณภาพข้อมูล) ----------
MIN_TURNOVER_M = 10.0   # มูลค่าซื้อขายเฉลี่ยขั้นต่ำ (ล้านบาท/วัน) ทั้งเฉลี่ย 10 วันและ 25 วัน
MIN_MCAP_B = 3.0        # มูลค่าตลาดขั้นต่ำ (พันล้านบาท)
MIN_FREEFLOAT = 15.0    # free float ขั้นต่ำ (%)
MIN_PRICE = 1.0         # กันหุ้นต่ำบาทที่ tick กระโดดแรงเป็น % มาก


def load_universe(df: pd.DataFrame) -> pd.DataFrame:
    u = df[
        (df['data_quality_ok'] == True)  # noqa: E712
        & (df['bars_stale'] == 0)
        & (df['ValMA25'] / 1e6 >= MIN_TURNOVER_M)
        & (df['avg_turnover_10d'] / 1e6 >= MIN_TURNOVER_M)
        & (df['Mcap'] / 1e9 >= MIN_MCAP_B)
        & (df['freefloatShares%'].fillna(0) >= MIN_FREEFLOAT)
        & (df['close'] >= MIN_PRICE)
    ].copy()
    u['vol_ratio'] = u['volume'] / u['VolMA25']
    u['ValMA25_M'] = u['ValMA25'] / 1e6
    u['Mcap_B'] = u['Mcap'] / 1e9
    u['ticker'] = u['symbol'].str.replace('SET:', '', regex=False)
    u['stop_2ATR'] = u['close_minus_2ATR']
    u['stop_dist%'] = (u['close'] - u['stop_2ATR']) / u['close'] * 100
    u['EMA26_price'] = (u['close'] / (1 + u['EMA26%C'] / 100)).round(2)
    return u


def setup_momentum(u: pd.DataFrame) -> pd.Series:
    return (
        (u['CDC'])                                  # EMA12 > EMA26 (CDC เขียว)
        & (u['EMA26%C'] > 0)                        # ราคายืนเหนือ EMA26
        & (u['EMA26%C'] <= 12)                      # แต่ไม่ยืดเกินไป
        & (u['MA25_60TF'])                          # MA25 > MA60
        & (u['RSIzone'] == 'zone3')                 # RSI วัน+สัปดาห์ > 50
        & (u['RSI14'].between(55, 78))              # โมเมนตัมดี ไม่สุดโต่ง
        & (u['MACD_hist'] > 0)                      # MACD เหนือ signal
        & (u['ADX'] >= 20) & (u['DI+'] > u['DI-'])  # เทรนด์มีแรง ฝั่งซื้อคุม
        & (u['RS_rank_20d'] >= 65)                  # ชนะตลาด (เปอร์เซ็นไทล์ 20 วัน)
        & (u['RS_rank_60d'] >= 50)
        & (u['ret_20d'] > 0)
        & (~u['RSI_div_bear'])                      # ไม่มีสัญญาณขัดแย้งขาลง
    )


def setup_breakout(u: pd.DataFrame) -> pd.Series:
    broke_level = (
        u['BreakMaxclose60'] | u['BreakPMax60']
        | u['BreakMaxclose100'] | u['BreakPMax100'] | u['Break_UC']
    )
    vol_confirm_strong = (u['vol_ratio'] >= 1.5) | u['OverVolMAx2'] | u['volmax25TF']
    fresh_signal = (
        (u['CDC_break'] | u['MACD_cross_up'] | u['EMA26_break'])
        & (u['vol_ratio'] >= 1.0)
        & (u['RSI14'] >= 50)
    )
    return (
        ((broke_level & vol_confirm_strong) | fresh_signal)
        & (u['EMA26%C'] > 0)
        & (u['ret_1d'] > 0)          # แท่งล่าสุดยังบวก ไม่ใช่ทะลุแล้วโดนทุบกลับ
        & (~u['RSI_div_bear'])
    )


def setup_pullback(u: pd.DataFrame) -> pd.Series:
    weekly_strong = (u['RSIzone'] == 'zone4') | (
        (u['RSIzone'] == 'zone3') & (u['ret_5d'] < 0) & (u['RSI14'] < 58)
    )
    return (
        weekly_strong
        & (u['close'] > u['MA100'])                 # เทรนด์ระยะยาวยังขึ้น
        & (u['close'] > u['MA60'])
        & (u['MA25_60TF'])
        & (u['EMA26%C'].between(-6, 3))             # ย่อมาใกล้ EMA26
        & (u['RSI14'] >= 35)                        # ไม่ใช่ขาลงรุนแรง
        & (u['ret_60d'] > 0)
        & (~u['RSI_div_bear'])
    )


def composite_score(u: pd.DataFrame) -> pd.Series:
    s = 0.25 * u['SCORE'].fillna(0)                       # คะแนน 8 เงื่อนไขจากไฟล์ (0-100)
    s += 0.20 * u['RS_rank_20d'].fillna(0)                # ความแข็งเทียบตลาด
    s += 0.10 * u['RS_rank_60d'].fillna(0)
    adx_pts = (u['ADX'].clip(20, 40) - 20) / 20 * 10      # ADX 20→40 = 0→10 คะแนน
    s += np.where(u['DI+'] > u['DI-'], adx_pts, 0)
    s += np.where(u['RSI14W'] > 55, 5, 0)                 # โมเมนตัมรายสัปดาห์
    s += np.where(u['MACD_hist_rising'], 5, 0)
    s += u['vol_ratio'].clip(0, 3) / 3 * 10               # วอลุ่มเข้า (เต็ม 10 ที่ 3 เท่า)
    s += np.where(u['MCDX_red'] > 5, 5, 0)                # แรงซื้อรายใหญ่ (MCDX)
    s += np.where(u['MCDX_red'] > u['MCDX_red-5'], 5, 0)  # และกำลังเพิ่มขึ้น
    s -= np.where(u['RSI_div_bear'], 10, 0)
    s -= np.where(u['EMA26%C'] > 10, 5, 0)                # ยืดไกลเส้นค่าเฉลี่ย เสี่ยงย่อ
    return s.round(1)


def warnings_col(u: pd.DataFrame) -> pd.Series:
    out = []
    for _, r in u.iterrows():
        msgs = []
        if pd.notna(r['days_to_earnings']) and 0 <= r['days_to_earnings'] <= 10:
            msgs.append(f"งบออกใน {int(r['days_to_earnings'])} วัน")
        if r['EMA26%C'] > 10:
            msgs.append('ราคายืดไกลเส้นค่าเฉลี่ย')
        if r['RSI14'] > 75:
            msgs.append('RSI สูง')
        if r['RSI_div_bear']:
            msgs.append('RSI divergence ขาลง')
        if r['beta'] > 1.5:
            msgs.append(f"beta สูง {r['beta']:.1f}")
        out.append(' | '.join(msgs))
    return pd.Series(out, index=u.index)


# ---------- ส่วนสร้างรายงาน Excel ----------
HDR = ['หุ้น', 'กลยุทธ์', 'ราคาปิด', 'คะแนนรวม', 'SCORE ไฟล์', 'RSI วัน', 'RSI สัปดาห์', 'ADX',
       'RS เทียบตลาด (pct)', 'ผลตอบแทน 20 วัน %', 'วอลุ่ม/เฉลี่ย 25 วัน (เท่า)',
       'แนวรับ EMA26', 'จุดตัดขาดทุน (2ATR)', 'ระยะถึงจุดตัด %', 'ATR ต่อวัน %',
       'มูลค่าซื้อขายเฉลี่ย 25 วัน (ลบ.)', 'มูลค่าตลาด (พันลบ.)', 'PE', 'ปันผล %',
       'กลุ่มธุรกิจ', 'คำเตือน']
ARIAL = 'Arial'


def _row_of(r):
    pe = round(r['PE'], 1) if pd.notna(r['PE']) else None
    dy = round(r['dividendYield'] * 100, 2) if pd.notna(r['dividendYield']) else None
    return [r['ticker'], r['setups'].strip(), r['close'], r['composite'], r['SCORE'],
            round(r['RSI14'], 1), round(r['RSI14W'], 1), round(r['ADX'], 1),
            round(r['RS_rank_20d'], 1), round(r['ret_20d'], 2), round(r['vol_ratio'], 2),
            r['EMA26_price'], round(r['stop_2ATR'], 2), round(r['stop_dist%'], 2),
            round(r['ATR%'], 2), round(r['ValMA25_M'], 1), round(r['Mcap_B'], 1),
            pe, dy, r['sector'], r['คำเตือน']]


def _write_sheet(ws, rows: pd.DataFrame):
    th_font = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
    th_fill = PatternFill('solid', fgColor='1F4E79')
    cell_font = Font(name=ARIAL, size=10)
    warn_fill = PatternFill('solid', fgColor='FFF2CC')
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = 'C2'
    for j, h in enumerate(HDR, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.border = th_font, th_fill, border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, (_, r) in enumerate(rows.iterrows(), 2):
        for j, v in enumerate(_row_of(r), 1):
            c = ws.cell(i, j, v)
            c.font, c.border = cell_font, border
            if isinstance(v, float):
                c.number_format = '#,##0.00'
        if r['คำเตือน']:
            for j in range(1, len(HDR) + 1):
                ws.cell(i, j).fill = warn_fill
    widths = [8, 22, 9, 9, 9, 8, 9, 8, 10, 10, 11, 10, 11, 9, 8, 12, 10, 7, 8, 22, 34]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 42


def build_report(u: pd.DataFrame, out_path: str, asof: str, set_close: float, set_chg20: float,
                 n_total: int):
    wb = Workbook()
    ws = wb.active
    ws.title = 'อ่านก่อน (เกณฑ์)'
    counts = {k: int(u['setups'].str.contains(k).sum())
              for k in ['A-Momentum', 'B-Breakout', 'C-Pullback']}
    info = [
        ('รายงานคัดกรองหุ้น SET สำหรับเทรด', ''),
        ('ข้อมูล ณ วันที่', asof),
        ('ภาวะตลาด', f'SET {set_close:,.2f} | ผลตอบแทนดัชนี 20 วัน {set_chg20:+.2f}%'),
        ('', ''),
        (f'Universe (ผ่าน {len(u)} ตัว จาก {n_total})',
         f'คุณภาพข้อมูลผ่าน + มูลค่าซื้อขายเฉลี่ย 10/25 วัน ≥ {MIN_TURNOVER_M:.0f} ลบ./วัน '
         f'+ มูลค่าตลาด ≥ {MIN_MCAP_B:.0f} พันลบ. + free float ≥ {MIN_FREEFLOAT:.0f}% + ราคา ≥ {MIN_PRICE:.0f} บาท'),
        (f"A-Momentum ({counts['A-Momentum']} ตัว)",
         'CDC เขียว (EMA12>EMA26), ราคาเหนือ EMA26 ไม่เกิน 12%, MA25>MA60, RSI วัน 55-78 และสัปดาห์ >50, '
         'MACD hist บวก, ADX≥20 และ DI+>DI-, RS rank 20 วัน ≥65, ไม่มี bearish divergence'),
        (f"B-Breakout ({counts['B-Breakout']} ตัว)",
         'ทะลุ high 60/100 วันหรือกรอบบน พร้อมวอลุ่ม ≥1.5 เท่าของเฉลี่ย 25 วัน '
         'หรือสัญญาณใหม่ (CDC/MACD/EMA26 เพิ่งตัดขึ้น) + วอลุ่มไม่ต่ำกว่าเฉลี่ย + แท่งล่าสุดบวก'),
        (f"C-Pullback ({counts['C-Pullback']} ตัว)",
         'ขาขึ้นรายสัปดาห์ยังแข็ง (RSI สัปดาห์ >50) แต่รายวันย่อลงมาใกล้ EMA26 (-6% ถึง +3%), '
         'ราคายังเหนือ MA60/MA100, ผลตอบแทน 60 วันเป็นบวก'),
        ('คะแนนรวม (composite)',
         'SCORE จากไฟล์ 25% + RS rank 20/60 วัน 30% + ADX 10 + RSI สัปดาห์ 5 + MACD เร่งขึ้น 5 '
         '+ วอลุ่มเข้า 10 + แรงซื้อรายใหญ่ MCDX 10 − โทษ divergence/ยืดตัว'),
        ('', ''),
        ('การใช้งาน',
         'จุดตัดขาดทุนอ้างอิง 2×ATR ใต้ราคาปิด หรือใต้แนวรับ EMA26 — เลือกอันที่ตื้นกว่าตามสไตล์ | '
         'ขนาดไม้: เสี่ยงต่อไม้ ~1-2% ของพอร์ต ÷ ระยะถึงจุดตัด'),
        ('คำเตือนสำคัญ',
         'หุ้นที่มี "งบออกใน N วัน" มีความเสี่ยงราคากระโดดข้ามคืนช่วงประกาศงบ ควรลดขนาดไม้หรือรอผ่านงบ'),
        ('ข้อจำกัด',
         'เป็นการคัดกรองทางเทคนิคจากข้อมูลในไฟล์เท่านั้น ไม่ใช่คำแนะนำการลงทุน ผลในอดีตไม่การันตีอนาคต '
         'โปรดพิจารณาปัจจัยพื้นฐาน/ข่าวประกอบและรับความเสี่ยงได้ก่อนเทรดจริง'),
    ]
    for i, (a, b) in enumerate(info, 1):
        ca, cb = ws.cell(i, 1, a), ws.cell(i, 2, b)
        ca.font = Font(name=ARIAL, bold=True, size=11 if i == 1 else 10)
        cb.font = Font(name=ARIAL, size=10)
        cb.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 150

    matched = u[u['setups'] != ''].sort_values('composite', ascending=False)
    _write_sheet(wb.create_sheet('Top 20 รวมทุกกลยุทธ์'), matched.head(20))
    for name in ['A-Momentum', 'B-Breakout', 'C-Pullback']:
        sub = u[u['setups'].str.contains(name)].sort_values('composite', ascending=False)
        _write_sheet(wb.create_sheet(name), sub)
    wb.save(out_path)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    df = pd.read_excel(src, sheet_name='ALL')
    asof = str(pd.to_datetime(df['datetime_'].iloc[0]).date())
    out = sys.argv[2] if len(sys.argv) > 2 else f'watchlist_{asof.replace("-", "")}.xlsx'

    u = load_universe(df)
    print(f'Universe หลังกรองสภาพคล่อง/คุณภาพ: {len(u)} ตัว จากทั้งหมด {len(df)}')

    u['composite'] = composite_score(u)
    u['คำเตือน'] = warnings_col(u)
    mA, mB, mC = setup_momentum(u), setup_breakout(u), setup_pullback(u)
    u['setups'] = ''
    u.loc[mA, 'setups'] += 'A-Momentum '
    u.loc[mB, 'setups'] += 'B-Breakout '
    u.loc[mC, 'setups'] += 'C-Pullback '

    show = ['ticker', 'close', 'composite', 'setups', 'RSI14', 'ADX', 'RS_rank_20d',
            'vol_ratio', 'ret_20d', 'stop_2ATR', 'ValMA25_M', 'sector', 'คำเตือน']
    for name, mask in [('A-Momentum', mA), ('B-Breakout', mB), ('C-Pullback', mC)]:
        sub = u[mask].sort_values('composite', ascending=False)
        print(f'\n===== {name}: {len(sub)} ตัว =====')
        print(sub[show].head(12).round(2).to_string(index=False))

    build_report(u, out, asof, df['SET_close'].iloc[0], df['SET_chg_20d'].iloc[0], len(df))
    print(f'\nบันทึกรายงานที่ {out}')


if __name__ == '__main__':
    main()
